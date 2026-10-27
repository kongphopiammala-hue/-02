import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Page Margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles Setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Angsana New'
    normal_style.font.size = Pt(15)
    normal_style.font.color.rgb = RGBColor(51, 65, 85) # Slate 700

    # Primary colors
    COLOR_PRIMARY = RGBColor(14, 116, 144)   # Cyan 700 / Deep Teal
    COLOR_SECONDARY = RGBColor(30, 41, 59)  # Slate 800
    COLOR_MUTED = RGBColor(100, 116, 139)    # Slate 500

    # --- TITLE / COVER BLOCK ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(20)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("คู่มือการใช้งานเว็บไซต์ร้านขายโมเดลและของสะสม")
    run_title.font.name = 'Angsana New'
    run_title.font.size = Pt(28)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(24)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("TITAN HOBBY & COLLECTIBLES (E-Commerce & Pre-order Platform)")
    run_sub.font.name = 'Angsana New'
    run_sub.font.size = Pt(18)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_SECONDARY

    # Metadata Box
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("ชื่อระบบ:", "TITAN HOBBY STORE - Collector Experience Platform"),
        ("เวอร์ชันเอกสาร:", "Version 1.0 (Demo Prototype)"),
        ("วันที่จัดทำ:", "27 ตุลาคม 2026"),
        ("สถานะระบบ:", "พร้อมใช้งาน (Local Server: http://localhost:8080)")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, 80, 80, 120, 120)
        set_cell_margins(c1, 80, 80, 120, 120)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_SECONDARY
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(v)
        r1.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # --- HELPER FUNCTIONS ---
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.font.name = 'Angsana New'
        r.font.size = Pt(20)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = 'Angsana New'
        r.font.size = Pt(17)
        r.font.bold = True
        r.font.color.rgb = COLOR_SECONDARY
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_SECONDARY
        r_txt = p.add_run(text)
        return p

    def add_callout(text, title="คำแนะนำเพิ่มเติม (TIP):"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        cell.width = Inches(6.5)
        set_cell_background(cell, "EFF6FF") # Soft blue
        set_cell_margins(cell, 120, 120, 160, 160)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r_t = p.add_run(title + " ")
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(29, 78, 216)
        r_c = p.add_run(text)
        r_c.font.color.rgb = RGBColor(30, 58, 138)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- SECTION 1 ---
    add_h1("1. บทนำและภาพรวมเว็บไซต์ (System Overview)")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.add_run("เว็บไซต์ ").font.color.rgb = RGBColor(51, 65, 85)
    r = p.add_run("TITAN HOBBY & COLLECTIBLES")
    r.font.bold = True
    p.add_run(" เป็นเว็บแอปพลิเคชันร้านค้าออนไลน์สำหรับจัดจำหน่ายและรับพรีออเดอร์โมเดล ฟิกเกอร์ Gunpla และ Art Toys ระดับพรีเมียม โดยได้รับการออกแบบขึ้นตามแนวคิด ")
    r_cc = p.add_run("“Collector-Centric Experience”")
    r_cc.font.bold = True
    p.add_run(" ที่เน้นตอบสนองพฤติกรรมและความต้องการของนักสะสมโดยเฉพาะ เช่น การหมุนดูรายละเอียดโมเดล 3 มิติ, การเปรียบเทียบขนาดจริงก่อนซื้อ, การเปิดจองแบบวางเงินมัดจำ และการจำลองระบบชำระเงินที่สะดวกรวดเร็ว")

    # Key Objectives Table
    tbl_obj = doc.add_table(rows=5, cols=2)
    tbl_obj.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = [("จุดเด่นของระบบ", "คำอธิบายความสำคัญต่อนักสะสม")]
    rows_data = [
        ("ธีม Cyber-Mecha Dark Mode", "คุมโทนสีดำ-เทาเข้มตัดด้วยไฟนีออน สื่อถึงบรรยากาศโรงเก็บหุ่นยนต์และห้องจัดแสดงโมเดล ขับให้ภาพโมเดลโดดเด่น"),
        ("Interactive 360° Turntable", "แท่นหมุนจำลอง 3 มิติ ให้นักสะสมคลิกลากหมุนดูองศาโมเดลได้ 360 องศาเพื่อตรวจดูรายละเอียดรอบตัว"),
        ("ระบบมัดจำ (Deposit System)", "แบ่งการชำระเงินเป็น มัดจำ หรือ จ่ายเต็มจำนวน ตอบโจทย์สินค้าราคาสูงและสินค้าพรีออเดอร์"),
        ("ตัววัดสเกล (Scale Visualizer)", "แสดงแถบเปรียบเทียบขนาดจริงกับสิ่งของทั่วไป (เช่น ขวดน้ำ, กระป๋อง) ช่วยประเมินพื้นที่ตู้โชว์")
    ]
    for row_idx, (c0_text, c1_text) in enumerate([headers[0]] + rows_data):
        row = tbl_obj.rows[row_idx]
        cell0, cell1 = row.cells[0], row.cells[1]
        cell0.width = Inches(2.3)
        cell1.width = Inches(4.2)
        set_cell_margins(cell0, 80, 80, 120, 120)
        set_cell_margins(cell1, 80, 80, 120, 120)
        if row_idx == 0:
            set_cell_background(cell0, "0E7490")
            set_cell_background(cell1, "0E7490")
            p0 = cell0.paragraphs[0]; r0 = p0.add_run(c0_text); r0.font.bold = True; r0.font.color.rgb = RGBColor(255, 255, 255)
            p1 = cell1.paragraphs[0]; r1 = p1.add_run(c1_text); r1.font.bold = True; r1.font.color.rgb = RGBColor(255, 255, 255)
        else:
            bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            set_cell_background(cell0, bg)
            set_cell_background(cell1, bg)
            p0 = cell0.paragraphs[0]; r0 = p0.add_run(c0_text); r0.font.bold = True; r0.font.color.rgb = COLOR_SECONDARY
            p1 = cell1.paragraphs[0]; r1 = p1.add_run(c1_text); r1.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- SECTION 2 ---
    add_h1("2. วิธีการเข้าใช้งานเว็บไซต์ (How to Access & Run)")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run("ผู้ใช้งานสามารถเข้าชมและทดสอบระบบได้ผ่าน 2 ช่องทางหลัก:")
    
    add_bullet("ช่องทางที่ 1 (แนะนำ): เข้าผ่านเว็บเบราว์เซอร์ทาง Local Server", "เปิดโปรแกรมเบราว์เซอร์ เช่น Google Chrome, Microsoft Edge หรือ Safari แล้วพิมพ์ URL: http://localhost:8080 จากนั้นกด Enter")
    add_bullet("ช่องทางที่ 2: เปิดไฟล์ตรงจากเครื่องคอมพิวเตอร์", "เข้าไปที่โฟลเดอร์ c:\\งาน02 แล้วดับเบิลคลิกที่ไฟล์ index.html เพื่อเปิดดูผ่านเบราว์เซอร์ได้ทันทีโดยไม่ต้องต่ออินเทอร์เน็ตเซิร์ฟเวอร์ภายนอก")

    add_callout("หากต้องการเริ่มเปิด Local Server ใหม่ในอนาคต ให้เปิด PowerShell ในโฟลเดอร์โปรเจกต์แล้วรันคำสั่ง: python -m http.server 8080", "หมายเหตุการรันเซิร์ฟเวอร์:")

    # --- SECTION 3 ---
    add_h1("3. คู่มือการใช้งานฟังก์ชันหลักบนเว็บไซต์ (Feature Walkthrough)")

    add_h2("3.1 แถบแจ้งเตือนและการค้นหาสินค้า (Navigation & Live Search)")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run("แถบด้านบนสุดถูกจัดวางเครื่องมือช่วยอำนวยความสะดวกในการค้นหาและเลือกหมวดหมู่:")
    add_bullet("แถบข้อความสิทธิประโยชน์ (Announcement Bar):", "แสดงข้อมูลความแท้ของสินค้า (100% Japan Lot) และแจ้งเตือนของแถมพิเศษประจำเดือน")
    add_bullet("ช่องค้นหาอัจฉริยะ (Live Search):", "สามารถพิมพ์คำค้นหา เช่น ชื่อหุ่น (Gundam, Asuka), ซีรีส์การ์ตูน, ค่ายผลิต (Bandai, Hobby Max) หรือขนาดสเกล (1/7, 1/60) ระบบจะกรองและแสดงผลการค้นหาให้ทันทีแบบเรียลไทม์โดยไม่ต้องกดปุ่มค้นหา")
    add_bullet("แถบหมวดหมู่ด่วน (Category Pills):", "คลิกปุ่มหมวดหมู่เพื่อดูเฉพาะกลุ่มสินค้าที่สนใจ เช่น Gunpla & Mecha, Scale Figures, Art Toys & Blind Box, สินค้าพรีออเดอร์ หรือสินค้าพร้อมส่ง")

    add_h2("3.2 แท่นหมุนโมเดล 3 มิติ 360 องศา (Interactive 360° Turntable)")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run("ส่วนจัดแสดงสินค้าไฮไลต์ (Hero Spotlight) บริเวณด้านบนของหน้าแรก มีจุดเด่นดังนี้:")
    add_bullet("การลากเมาส์เพื่อหมุน (Drag to Rotate):", "คลิกเมาส์ค้างไว้บนแท่นจัดแสดง แล้วลากไปทางซ้ายหรือขวาเพื่อหมุนดูตัวโมเดลได้ 360 องศา พร้อมทั้งมีตัวเลขระบุองศา (Degree Indicator) แจ้งให้ทราบแบบเรียลไทม์ (รองรับการใช้นิ้วปัดบนมือถือหรือแท็บเล็ต)")
    add_bullet("ปุ่มตั้งค่ามุมมองลัด (Preset Angle Buttons):", "คลิกปุ่ม 0° (หน้าตรง), 90° (ด้านข้าง), 180° (ด้านหลัง) หรือ 270° (ด้านขวา) เพื่อให้โมเดลหมุนไปยังมุมที่ต้องการได้ทันที")
    add_bullet("นาฬิกานับถอยหลังรอบปิดรับจอง (Countdown Timer):", "แสดงเวลาที่เหลือก่อนปิดรับจองแบบ วัน:ชั่วโมง:นาที:วินาที เพื่อช่วยให้นักสะสมไม่พลาดรอบสำคัญ")
    add_bullet("ปุ่มจองล่วงหน้าด่วน (Quick Pre-order):", "คลิกปุ่ม 'จองล่วงหน้าทันที' ระบบจะหยิบสินค้าพร้อมตั้งค่ายอดมัดจำลงในตระกร้าและเปิดลิ้นชักตระกร้าให้ทันทีในคลิกเดียว")

    add_h2("3.3 แคตตาล็อกและการกรองสินค้าละเอียด (Catalog & Advanced Filters)")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run("ในส่วนของรายการสินค้าหลัก ผู้ใช้สามารถใช้ตัวกรองเฉพาะทางสำหรับนักสะสมได้ดังนี้:")
    add_bullet("ตัวกรองสถานะ (Status Filter):", "เลือกดูเฉพาะ 'สินค้าพร้อมส่งทันที', 'สินค้าเปิดรับพรีออเดอร์' หรือ 'สินค้า Limited Edition'")
    add_bullet("ตัวกรองขนาดสเกล (Scale Filter):", "เลือกคัดกรองตามสเกล เช่น 1/60 (PG), 1/100 (MG), 1/144 (RG/HG), 1/7 Scale, 1/6 Scale หรือ Non-scale สำหรับกล่องสุ่ม")
    add_bullet("การจัดเรียง (Sorting):", "จัดเรียงตามยอดนิยม, ราคาจากน้อยไปมาก, ราคาจากมากไปน้อย หรือสินค้าที่เปิดพรีออเดอร์ล่าสุด")
    add_bullet("ปุ่มล้างตัวกรอง (Reset Filter):", "สามารถกดปุ่มรีเซ็ตเพื่อคืนค่าตัวกรองทั้งหมดกลับเป็นค่าเริ่มต้นได้ในทันที")

    add_h2("3.4 หน้าต่างดูรายละเอียดเชิงลึกและเครื่องมือวัดสเกล (Product Modal & Scale Visualizer)")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run("เมื่อคลิกที่การ์ดสินค้าหรือปุ่ม 'ดูรายละเอียด & สเกล' ระบบจะเปิดหน้าต่างป๊อปอัปขึ้นมา ประกอบไปด้วย:")
    add_bullet("แกลเลอรีภาพหลายมุมมอง:", "สามารถคลิกภาพย่อย (Thumbnails) ด้านล่างเพื่อดูภาพมุมต่างๆ ของชิ้นงาน เช่น ท่าแอ็กชัน, รอยดีคอล หรือแพ็กเกจกล่อง")
    add_bullet("แถบเปรียบเทียบสเกลจริง (Scale Comparison Bar):", "แสดงกราฟฟิกเปรียบเทียบความสูงจริงของโมเดลเทียบกับขวดน้ำดื่ม (30 cm) และกระป๋องเครื่องดื่ม (12 cm) เพื่อให้นักสะสมกะขนาดและเตรียมพื้นที่วางในตู้โชว์ได้อย่างแม่นยำ")
    add_bullet("ตารางสเปกสินค้า:", "ระบุค่ายผู้ผลิต, ซีรีส์อนิเมะ, วัสดุที่ใช้ผลิต (PVC, ABS, Diecast), ล็อตตัวแทนจำหน่าย และกำหนดการสินค้าถึงประเทศไทย")
    add_bullet("ตัวเลือกรูปแบบชำระเงิน:", "สามารถเลือกว่าต้องการ 'จ่ายมัดจำ' (สำหรับสินค้าพรีออเดอร์) หรือ 'จ่ายเต็มจำนวน' (เพื่อรับสิทธิ์ส่งฟรีทันที)")

    add_h2("3.5 ระบบตระกร้าสินค้า (Shopping Cart Drawer)")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run("ลิ้นชักตระกร้าสินค้าจะเปิดออกมาจากด้านขวาของหน้าจอ มีการแยกรายละเอียดการคิดเงินอย่างชัดเจน:")
    add_bullet("แสดงประเภทการชำระต่อชิ้น:", "ระบุชัดเจนว่ารายการใดเป็นการสั่งแบบมัดจำ หรือจ่ายเต็มจำนวน")
    add_bullet("ปรับเพิ่ม/ลดจำนวนสินค้า:", "กดปุ่ม (+) หรือ (-) เพื่อปรับจำนวน หรือกดไอคอนถังขยะเพื่อลบรายการออกจากตระกร้า")
    add_bullet("สรุปยอด 2 ระดับ:", "ระบบจะคำนวณ 'ยอดรวมมูลค่าสินค้าทั้งหมด' และ 'ยอดเงินที่ต้องชำระทันทีในปัจจุบัน (Due Now)' ให้เห็นคู่กันอย่างโปร่งใส")

    add_h2("3.6 ขั้นตอนการสั่งซื้อและจำลองการจ่ายเงิน (Checkout Gateway & QR PromptPay)")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run("เมื่อกดปุ่ม 'ดำเนินการสั่งซื้อ & ชำระเงิน' ระบบจะนำเข้าสู่ขั้นตอนชำระเงินแบบ 3 สเต็ป:")
    add_bullet("สเต็ปที่ 1 (ข้อมูลผู้รับ):", "กรอกชื่อ-นามสกุล, เบอร์โทรศัพท์, LINE ID (สำหรับรับข้อความแจ้งเตือนเมื่อโมเดลเข้าไทย) และที่อยู่จัดส่งพัสดุ")
    add_bullet("สเต็ปที่ 2 (ชำระเงิน PromptPay):", "ระบบจะสร้าง QR Code พร้อมยอดเงินที่ต้องชำระขึ้นมาจำลองการสแกนจ่ายผ่านโมบายแบงก์กิ้ง พร้อมปุ่มทดสอบยืนยันการโอน")
    add_bullet("สเต็ปที่ 3 (ยืนยันคำสั่งซื้อ):", "แสดงรหัสคำสั่งซื้อ (Order Tracking ID เช่น #TH-2026-XXXX) พร้อมแจ้งสถานะการส่งใบเสร็จและใบจองไปยังผู้ใช้งาน")

    add_h2("3.7 โซนชุมชนนักสะสม (Collector's Hall of Fame)")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run("อยู่บริเวณท้ายหน้าแรก เป็นโซนรวมรูปถ่ายรีวิวโมเดลจากตู้โชว์ของลูกค้าจริงทั่วประเทศ ช่วยเพิ่มความมั่นใจด้านการแพ็กห่อกันกระแทก ความคมของมุมกล่อง และการันตีของแท้ 100%")

    # --- SECTION 4 ---
    add_h1("4. คำถามที่พบบ่อยและการปรับแต่งโค้ด (FAQ & Customization Guide)")
    
    faq_table = doc.add_table(rows=4, cols=2)
    faq_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    faq_data = [
        ("คำถามที่พบบ่อย", "แนวทางแก้ไข / คำอธิบาย"),
        ("ต้องการเพิ่มสินค้าตัวใหม่ลงในเว็บทำอย่างไร?", "เปิดไฟล์ app.js แล้วไปที่อาร์เรย์ PRODUCTS จากนั้นคัดลอกอ็อบเจกต์สินค้าแล้วใส่ข้อมูล เช่น ชื่อ, ราคา, สเกล และลิงก์รูปภาพ"),
        ("รูปภาพสินค้าสามารถใช้รูปจากเครื่องตัวเองได้ไหม?", "สามารถทำได้ โดยสร้างโฟลเดอร์ images ใน c:\\งาน02 แล้วใส่ชื่อไฟล์ภาพในฟิลด์ images ของไฟล์ app.js เช่น 'images/my-gundam.jpg'"),
        ("ต้องการเปลี่ยนสีธีมหรือสีไฟนีออน?", "สามารถปรับเปลี่ยนค่าสีในไฟล์ index.html ตรงส่วน tailwind.config หรือแก้สไตล์เอฟเฟกต์แสงเงาในไฟล์ style.css")
    ]
    for row_idx, (c0_text, c1_text) in enumerate(faq_data):
        row = faq_table.rows[row_idx]
        cell0, cell1 = row.cells[0], row.cells[1]
        cell0.width = Inches(2.3)
        cell1.width = Inches(4.2)
        set_cell_margins(cell0, 80, 80, 120, 120)
        set_cell_margins(cell1, 80, 80, 120, 120)
        if row_idx == 0:
            set_cell_background(cell0, "1E293B")
            set_cell_background(cell1, "1E293B")
            p0 = cell0.paragraphs[0]; r0 = p0.add_run(c0_text); r0.font.bold = True; r0.font.color.rgb = RGBColor(255, 255, 255)
            p1 = cell1.paragraphs[0]; r1 = p1.add_run(c1_text); r1.font.bold = True; r1.font.color.rgb = RGBColor(255, 255, 255)
        else:
            bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            set_cell_background(cell0, bg)
            set_cell_background(cell1, bg)
            p0 = cell0.paragraphs[0]; r0 = p0.add_run(c0_text); r0.font.bold = True; r0.font.color.rgb = COLOR_SECONDARY
            p1 = cell1.paragraphs[0]; r1 = p1.add_run(c1_text); r1.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    # Footer note
    foot_p = doc.add_paragraph()
    foot_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = foot_p.add_run("— จัดทำโดย Antigravity AI เพื่อโครงการร้านค้าโมเดล TITAN HOBBY & COLLECTIBLES —")
    r_foot.font.italic = True
    r_foot.font.size = Pt(13)
    r_foot.font.color.rgb = COLOR_MUTED

    # Save to c:\งาน02
    output_path = r"c:\งาน02\คู่มือการใช้งานเว็บไซต์_TITAN_HOBBY.docx"
    doc.save(output_path)
    print("Document successfully created at:", output_path)

if __name__ == '__main__':
    create_document()
